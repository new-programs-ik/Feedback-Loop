"""Build "The Formula Search - The Simple Version.docx": what we did when we were asked to throw
away the given formula and find the real one, written the way a PM would explain it to a manager.

Same house style as build_sentiment_simple_doc.py. Every number comes from the study's own output
(analysis/out/formula_results.json, formula_trust.json) or from Formula-One-Pager.pdf.

    python analysis/build_formula_simple_doc.py
"""
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.environ.get("FORMULA_DOC_OUT") or os.path.join(ROOT, "The Formula Search - The Simple Version.docx")

INK = RGBColor(0x16, 0x16, 0x1A)
MUTED = RGBColor(0x4B, 0x4B, 0x53)
BRAND = RGBColor(0x2A, 0x78, 0xD6)
GOOD = RGBColor(0x1B, 0x7F, 0x5E)
BAD = RGBColor(0xB4, 0x3A, 0x2E)
HEADER_FILL = "EEF3FB"
ZEBRA_FILL = "F7F7F5"
NOTE_FILL = "F3F7FD"

doc = Document()
sec = doc.sections[0]
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin = sec.bottom_margin = Cm(1.8)

base = doc.styles["Normal"]
base.font.name = "Calibri"
base.font.size = Pt(11)
base.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
for name, size in (("Heading 1", 15), ("Heading 2", 12.5)):
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = INK
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def para(text="", size=11, bold=False, color=None, italic=False, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(text, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet")
    if bold_lead:
        r = p.add_run(bold_lead + " ")
        r.bold = True
    p.add_run(text)
    p.paragraph_format.space_after = Pt(4)
    return p


def note(title, body):
    """A one-cell tinted box for the 'what this number means' asides."""
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = t.rows[0].cells[0]
    cell.text = ""
    p1 = cell.paragraphs[0]
    r = p1.add_run(title)
    r.bold = True
    r.font.size = Pt(10.5)
    p2 = cell.add_paragraph()
    r2 = p2.add_run(body)
    r2.font.size = Pt(10.5)
    p2.paragraph_format.space_after = Pt(2)
    shade(cell, NOTE_FILL)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def table(headers, rows, widths=None, bold_first_col=False, colors=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        r = hdr[i].paragraphs[0].add_run(h)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = MUTED
        shade(hdr[i], HEADER_FILL)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(10.5)
            if bold_first_col and i == 0:
                r.bold = True
            if colors and (ri, i) in colors:
                r.font.color.rgb = colors[(ri, i)]
                r.bold = True
            if ri % 2 == 1:
                shade(cells[i], ZEBRA_FILL)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


# ─────────────────────────────────────────────────────────────────── title
para("INTERVIEW KICKSTART · NEW PROGRAMS", size=9, bold=True, color=BRAND, space_after=2)
para("Did we find a better formula? — the simple version", size=20, bold=True, space_after=2)
para("We were told to forget the formula we were given and go find the real one. This is what we did, "
     "what we found, and what the team has to decide. Nothing here needs maths.",
     size=12, color=MUTED, space_after=10)

# ─────────────────────────────────────────────────────────────────── 1
doc.add_heading("1. The question", level=1)
para("Every class gets one score out of 100, and the score puts it in a band: Excellent, Good, Average "
     "or Bad. The band decides the work — Bad gets a video analysis, Average gets a transcript read, "
     "Good and Excellent are left alone.")
para("The formula behind that score came from one person's judgement. Nobody had checked whether it is "
     "the right formula — whether those are the right things to measure, in the right proportions, in "
     "the right shape. So the question was simple:")
para("“Is this the best formula we can have — and if not, what is?”", bold=True, space_after=10)

# ─────────────────────────────────────────────────────────────────── 2
doc.add_heading("2. How we looked", level=1)
para("Six steps, run over and over in a loop until nothing improved. The loop ran 28 times.", color=MUTED)
table(
    ["Step", "What we did", "Why"],
    [
        ["1", "Listed everything we can possibly know about a class at the moment we score it — 43 things: "
              "the rating, how many rated, how many attended, the vote, the instructor's past classes, how "
              "the same module went for other cohorts, where the class sits in the course, what the cohort "
              "did last week, and so on.",
         "You cannot pick the right ingredients before you know every ingredient available."],
        ["2", "Checked which of those 43 actually warn us that something is going wrong.",
         "Anything that does not warn us has no business being in the formula."],
        ["3", "Built five different kinds of formula — not only the one we had.",
         "The shape of a formula matters as much as its ingredients."],
        ["4", "Invented 378,046 imaginary classes covering every combination we could construct — including "
              "situations that have never happened here yet.",
         "Bishal's point: a case that has not happened yet will happen one day. The formula must not "
         "break when it does."],
        ["5", "Wrote 23 rules of good behaviour a formula must never break, and tested every formula "
              "against every rule.",
         "For example: more “no” votes must never raise a score. Obvious to a human, easy for a formula "
         "to get wrong."],
        ["6", "Hid the last three months of real classes. Built and tuned everything on January–May, then "
              "tested on the hidden June–August.",
         "Any formula can be made to look good on data it has already seen. Only hidden months are "
         "honest."],
    ],
    widths=[1.2, 8.6, 7.2],
)

# ─────────────────────────────────────────────────────────────────── 3
doc.add_heading("3. What we tried", level=1)
table(
    ["The formula", "In plain words", "What happened"],
    [
        ["The manager's original", "Rating out of 60, 30 points if 80% want the instructor back, 6 points "
                                   "if ten or more rated, 4 points for turnout.",
         "Hides 105 low-rated classes behind a Good or Excellent badge. Breaks 8 of the 23 rules."],
        ["Today's version", "Rating out of 60 (steeper below 4.55), the vote out of 25 (gradual, not "
                            "all-or-nothing), the instructor's record out of 15, plus the two agreed lines "
                            "as hard limits.",
         "The winner. Nothing beat it, and it is the only one that keeps all 23 rules."],
        ["Today's version + a trust guard", "The same, but a class with very few raters is blended with "
                                            "about three typical classes of that module first.",
         "Same verdicts, steadier number. Needs one change in the code before it can be switched on. "
         "Kept in the drawer."],
        ["Add more ingredients", "Add the module's own history, what the cohort did last week, whether "
                                 "attendance dropped.",
         "Each one adds nothing once the months are hidden. Dropped."],
        ["A statistician's version of the vote", "Instead of the plain percentage, use the cautious "
                                                 "lower estimate that statisticians use for small samples.",
         "Treats a thin vote as if it were a bad vote. Breaks the rules. Dropped."],
        ["A probability model", "Ignore points entirely; let a model output “the chance this class needs "
                                "attention” and turn that into 0–100.",
         "Nine classes in ten land between 54 and 86, so the bands lose their meaning — and no PM can "
         "work it out by hand. Dropped."],
        ["The instructor's series", "Score the instructor's run of classes rather than the class.",
         "Counts the instructor's record twice, so one bad history sinks a genuinely fine class. Dropped."],
    ],
    widths=[4.4, 6.6, 6.0], bold_first_col=True,
)

# ─────────────────────────────────────────────────────────────────── 4
doc.add_heading("4. The result", level=1)
para("The formula we already use is the best one we could find. That is a real result, not a shrug: it "
     "now stands on eight months of evidence instead of one person's judgement, and we know exactly how "
     "much better it is than the original.", space_after=8)
para("Tested only on the three hidden months — 1,027 classes the formula had never seen:", color=MUTED)
table(
    ["What we measured", "The manager's original", "Today's version"],
    [
        ["Low-rated classes wrongly shown as Good or Excellent", "105 of 873", "0"],
        ["Classes marked “Bad” that were actually rated 4.55 or better", "12 of 95", "0"],
        ["Verdict changes if one learner votes differently", "26 in 100", "19 in 100"],
        ["Warns us the instructor's next class will go wrong", "0.65", "0.67"],
        ["Chance the next class goes wrong: Bad band vs Excellent band", "45% vs 14%", "54% vs 10%"],
        ["Rules of good behaviour kept, out of 23", "15", "23"],
    ],
    widths=[8.4, 4.3, 4.3], bold_first_col=True,
    colors={(0, 1): BAD, (0, 2): GOOD, (1, 1): BAD, (1, 2): GOOD, (5, 1): BAD, (5, 2): GOOD},
)
note("What that 0.65 and 0.67 mean",
     "Take one class that later went wrong and one that did not, at random. How often does the score put "
     "the wrong one lower? A coin toss would score 0.50. The original manages 0.65, today's version 0.67. "
     "So the score is genuinely useful — but it is a warning light, not a crystal ball, and no formula we "
     "tried does much better on this data.")

# ─────────────────────────────────────────────────────────────────── 5
doc.add_heading("5. Small classes and unfair raters — what the data says", level=1)
para("Bishal's point: if only three people rate a class and two of them mark it down unfairly, the whole "
     "class looks bad. So we measured exactly when a low rating can be believed. We asked: when a class "
     "rated below 4.55, how often did that cohort's next class also go wrong?", space_after=8)
table(
    ["Learners who rated", "Class was rated low", "Class was fine", "What it tells us"],
    [
        ["1 or 2", "23% went wrong next", "20% went wrong next", "Nothing. A low rating from one or two "
                                                                 "people means nothing at all."],
        ["3 or 4", "33%", "16%", "It starts to matter — twice the risk."],
        ["5 to 9", "35%", "15%", "Believable."],
        ["10 to 14", "31%", "13%", "Believable."],
        ["15 or more", "38%", "17%", "Believable."],
    ],
    widths=[3.4, 4.0, 3.4, 6.2], bold_first_col=True,
)
para("So the app's rules match the evidence: no band under 3 votes, and no class is sent for analysis "
     "under 6 votes — it is watched instead. Two or three unhappy learners can never put a class in the "
     "queue on their own.", bold=True)
bullet("Head-count should never add or remove points — it should only decide how much we believe the "
       "rating. The old “6 points if ten or more rated” did the opposite: it rewarded big rooms and "
       "punished test reviews, which are smaller by nature.",
       bold_lead="What we changed our mind about.")
bullet("We expected a low rating from few raters to be more suspicious than a high one. It is not — both "
       "are simply uncertain, and the data shows no difference.",
       bold_lead="What surprised us.")
bullet("Blending a thin class with about three typical classes of the same module steadies the number "
       "without changing any verdict. Ready to switch on after one code change.",
       bold_lead="What is waiting.")
bullet("The real fix is knowing each rater's own history — does this learner mark everything down? That "
       "needs learner-level ratings, which we do not have yet. The slot for it is designed.",
       bold_lead="What we still cannot do.")

# ─────────────────────────────────────────────────────────────────── 6
doc.add_heading("6. Three things that are not the formula's fault", level=1)
para("These are decisions for the team. No formula can settle them.", color=MUTED)
table(
    ["The problem", "What is happening", "The choice"],
    [
        ["The queue is bigger than the team",
         "June–August produced about 14 analyses a week; the team can do about 12. Classes genuinely got "
         "worse over the year: 16% were under a line in January, 24% in August.",
         "Do more analyses, or only analyse classes with 7+ votes (that gives 13 a week), or move the "
         "agreed lines."],
        ["One vote can change the verdict",
         "One class in five sits within a single vote of the 4.55 rating line or the 80% approval line. "
         "So one learner changing their mind flips it.",
         "Accept it, or turn the hard lines into a narrow grey zone where a class is watched instead of "
         "flipped."],
        ["We are judging by proxy",
         "Without learner-level data we cannot see “a learner was let down”. We infer it from what "
         "happens next. And PMs' own decisions are not recorded — one confirm in the whole database.",
         "Ask every PM to press Confirm or Dismiss on each flagged class. In three months that alone "
         "would sharpen the next study more than any formula change."],
    ],
    widths=[4.4, 7.0, 5.6], bold_first_col=True,
)

# ─────────────────────────────────────────────────────────────────── 7
doc.add_heading("7. If you tell your manager one thing", level=1)
note("The paragraph to send",
     "We were asked to check whether the class score's formula is the right one, so we rebuilt it from "
     "scratch: 43 possible ingredients, five kinds of formula, 28 rounds of trial and error, 378,046 "
     "invented test cases, and 23 written rules of good behaviour — all tested on three months of real "
     "classes that were deliberately hidden from the work. The formula we are already using came out on "
     "top: it is the only one that keeps every rule, it never hides a low-rated class behind a good "
     "badge (the original hid 105), and it warns us about the next class better than the original does. "
     "So we are keeping it — now backed by eight months of evidence rather than judgement. What the study "
     "does ask the team to decide is not the formula but the workload: the queue currently runs at about "
     "14 analyses a week against a capacity of 12, because classes genuinely got worse over the year.")

para("")
para("Full detail, with every chart and number: Formula-Study.pdf (22 pages) and Formula-One-Pager.pdf. "
     "The study can be re-run with one command whenever new data arrives — we plan to re-run it monthly.",
     size=9, color=MUTED, italic=True)

doc.save(OUT)
print("wrote", OUT)
